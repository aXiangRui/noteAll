"""Phase 0：PPT → 复习资料（大纲 / 知识点 / 考核重点）→ HTML → PDF

只用于**验证输出质量**，不进产品。产品里这三步是 C++ 的 Tool 节点。

用法：
  python phase0.py <课件.pptx> <输出目录>

流程：
  1. 抽取（python-pptx）：每页 标题 / 要点 / 备注
  2. 生成（DeepSeek）：
     - 每 3 页一次「抽要点」（map）
     - 全部合起来一次「大纲 + 知识点 + 考核重点」（reduce）
  3. 渲染：HTML + CSS（改 CSS 就换风格）
  4. 出 PDF：Edge headless --print-to-pdf

关键约束（对应"考核重点不能凭空生成"）：
  * 每条 examFocus 必须带 pages（页码）和 basis（依据），依据只认：
    PPT 标记（★/必考/重点/例题）、老师强调（备注里的话）、高频出现
  * 没有依据的一律不许输出
"""
import json
import os
import re
import subprocess
import sys
import urllib.request
from pathlib import Path

from pptx import Presentation

API_BASE = os.environ.get("NOTE_LLM_BASE", "https://api.deepseek.com")
API_PATH = os.environ.get("NOTE_LLM_PATH", "/chat/completions")
API_MODEL = os.environ.get("NOTE_LLM_MODEL", "deepseek-chat")


# --------------------------------------------------------------------------- #
# 1. 抽取
# --------------------------------------------------------------------------- #
def extract(pptx_path: Path) -> dict:
    prs = Presentation(str(pptx_path))
    slides = []

    for index, slide in enumerate(prs.slides, start=1):
        title = ""
        bullets = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            is_title = shape == slide.shapes.title
            lines = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
            if is_title:
                title = lines[0] if lines else ""
            else:
                bullets.extend(lines)

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        slides.append({"page": index, "title": title, "bullets": bullets, "notes": notes})

    return {"source": pptx_path.name, "slides": slides}


# --------------------------------------------------------------------------- #
# 2. 生成
# --------------------------------------------------------------------------- #
def read_key() -> str:
    for candidate in ("dist/note.key", "dist/note_gui/note.key"):
        path = Path(candidate)
        if path.exists():
            return path.read_text(encoding="utf-8").strip()
    key = os.environ.get("NOTE_LLM_KEY")
    if key:
        return key.strip()
    raise SystemExit("找不到 API key（dist/note.key 或 NOTE_LLM_KEY）")


def chat(key: str, system: str, user: str, max_tokens: int = 2000) -> str:
    body = json.dumps(
        {
            "model": API_MODEL,
            "messages": [
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
            "temperature": 0.2,
            "max_tokens": max_tokens,
        }
    ).encode("utf-8")

    request = urllib.request.Request(
        API_BASE + API_PATH,
        data=body,
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {key}"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=180) as response:
        payload = json.loads(response.read().decode("utf-8"))

    return payload["choices"][0]["message"]["content"]


def json_from_text(text: str) -> dict:
    """模型偶尔会在 JSON 外面包 ```json ... ```，这里兜一下。"""
    fenced = re.search(r"```(?:json)?\s*(.+?)```", text, re.S)
    if fenced:
        text = fenced.group(1)
    start = text.find("{")
    end = text.rfind("}")
    if start >= 0 and end > start:
        text = text[start : end + 1]
    return json.loads(text)


MAP_SYSTEM = """你是课程助教，负责把课件片段整理成结构化要点。
只输出 JSON，不要解释。禁止编造：只写课件里确实出现的概念。
JSON 形状：{"segments":[{"pages":[页码],"points":[{"text":"要点，20 字以内","pages":[页码]}]}]}"""

REDUCE_SYSTEM = """你是课程助教，负责把多段要点整合成一份复习资料。
只输出 JSON，不要解释。

硬性规则：
1. 每条 examFocus 必须带 pages（页码数组）和 basis（依据），basis 只能是：
   "PPT 标记"（课件里有 ★ 必考 重点 例题 等字样）、
   "老师强调"（演讲者备注里的话）、
   "高频出现"（同一概念在多个片段重复出现）
2. 找不到依据就不要写这条 examFocus，宁可少写
3. 每条 point 也要带 pages

结构要求：**按主题合并成 3–5 章**（一页一章是错的）。
每章覆盖连续的页范围，pages 写 [起始页, 结束页]。

依据（basis）判定：
  "老师强调"  ← 演讲者备注里出现的话（备注是老师口头补充的）
  "PPT 标记"  ← 幻灯片正文里出现 ★ 必考 重点 例题 等字样
  "高频出现"  ← 同一概念在多个片段重复出现

JSON 形状：
{"chapters":[{"title":"...","pages":[起,止],"outline":["..."],
  "points":[{"text":"...","pages":[...]}],
  "examFocus":[{"text":"...","pages":[...],"basis":"PPT 标记|老师强调|高频出现"}]}],
 "examSummary":["全篇最重要的 3-5 条，必须能在上面各章 examFocus 中找到对应"]}"""


def generate(key: str, extracted: dict, chunk_size: int = 3) -> dict:
    slides = extracted["slides"]

    # ---- map：每 chunk_size 页抽一次要点 ----
    segments = []
    for start in range(0, len(slides), chunk_size):
        chunk = slides[start : start + chunk_size]
        text = "\n\n".join(
            f"第 {s['page']} 页\n标题：{s['title']}\n要点：\n" + "\n".join(f"- {b}" for b in s["bullets"])
            + (f"\n备注：{s['notes']}" if s["notes"] else "")
            for s in chunk
        )
        reply = chat(key, MAP_SYSTEM, f"课件片段：\n{text}", max_tokens=1200)
        segments.extend(json_from_text(reply)["segments"])
        print(f"  map: 第 {start + 1}-{start + len(chunk)} 页 → {len(segments[-1]['points'])} 条要点")

    # ---- reduce：整合成复习资料 ----
    merged = json.dumps({"segments": segments}, ensure_ascii=False)
    reply = chat(key, REDUCE_SYSTEM, f"片段要点：\n{merged}", max_tokens=3000)
    return json_from_text(reply)


# --------------------------------------------------------------------------- #
# 3. 渲染
# --------------------------------------------------------------------------- #
HTML_TEMPLATE = """<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>{title} 复习资料</title>
<style>
  @page {{ size: A4; margin: 18mm 16mm; }}
  body {{ font-family: "Source Han Sans SC", "Microsoft YaHei", sans-serif;
         font-size: 11pt; line-height: 1.7; color: #1c1f23; }}
  h1 {{ font-size: 20pt; margin: 0 0 4pt; }}
  .sub {{ color: #6b7280; font-size: 9.5pt; margin-bottom: 14pt; }}
  h2 {{ font-size: 14pt; margin: 18pt 0 6pt; padding-bottom: 3pt;
        border-bottom: 1.5pt solid #2f6fd0; color: #23406e; }}
  h3 {{ font-size: 11.5pt; margin: 12pt 0 4pt; color: #2f6fd0; }}
  ul {{ margin: 4pt 0 4pt 16pt; padding: 0; }}
  li {{ margin: 2pt 0; }}
  .page {{ color: #9aa3ad; font-size: 8.5pt; margin-left: 4pt; }}
  .focus {{ background: #fff8e6; border-left: 3pt solid #e0a338;
            padding: 6pt 10pt; margin: 6pt 0; }}
  .basis {{ color: #a15c00; font-size: 8.5pt; }}
  .summary {{ background: #eef4ff; border-left: 3pt solid #2f6fd0;
              padding: 8pt 12pt; margin: 10pt 0 0; }}
  @media print {{ h2 {{ page-break-after: avoid; }} .chapter {{ page-break-inside: avoid; }} }}
</style></head><body>
<h1>{title}</h1>
<div class="sub">{source} · 共 {pages} 页 · 由课件自动整理，请对照页码核对</div>
{body}
</body></html>
"""


def render(extracted: dict, material: dict, out_html: Path) -> None:
    parts = []

    for chapter in material.get("chapters", []):
        pages = chapter.get("pages") or []
        span = f"第 {pages[0]}–{pages[-1]} 页" if len(pages) >= 2 else ""
        parts.append(f'<div class="chapter"><h2>{chapter["title"]} <span class="page">{span}</span></h2>')

        if chapter.get("outline"):
            parts.append("<h3>大纲</h3><ul>")
            parts.extend(f"<li>{item}</li>" for item in chapter["outline"])
            parts.append("</ul>")

        if chapter.get("points"):
            parts.append("<h3>知识点</h3><ul>")
            for point in chapter["points"]:
                refs = "".join(f'<span class="page">p{page}</span>' for page in point.get("pages", []))
                parts.append(f'<li>{point["text"]}{refs}</li>')
            parts.append("</ul>")

        if chapter.get("examFocus"):
            parts.append("<h3>考核重点</h3>")
            for focus in chapter["examFocus"]:
                refs = "".join(f'<span class="page">p{page}</span>' for page in focus.get("pages", []))
                basis = focus.get("basis", "")
                parts.append(
                    f'<div class="focus">{focus["text"]}{refs} '
                    f'<span class="basis">依据：{basis}</span></div>'
                )

        parts.append("</div>")

    if material.get("examSummary"):
        parts.append('<div class="summary"><h3>全篇重点</h3><ul>')
        parts.extend(f"<li>{item}</li>" for item in material["examSummary"])
        parts.append("</ul></div>")

    html = HTML_TEMPLATE.format(
        title=Path(extracted["source"]).stem,
        source=extracted["source"],
        pages=len(extracted["slides"]),
        body="\n".join(parts),
    )
    out_html.write_text(html, encoding="utf-8")


# --------------------------------------------------------------------------- #
# 4. PDF
# --------------------------------------------------------------------------- #
def to_pdf(html_path: Path, pdf_path: Path) -> None:
    edge = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
    subprocess.run(
        [
            edge,
            "--headless",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_path.resolve()}",
            html_path.resolve().as_uri(),
        ],
        check=True,
        timeout=120,
        capture_output=True,
    )


def main() -> None:
    if len(sys.argv) < 3:
        raise SystemExit("用法：python phase0.py <课件.pptx> <输出目录>")

    pptx_path = Path(sys.argv[1])
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)

    print("[1/4] 抽取")
    extracted = extract(pptx_path)
    (out_dir / "extracted.json").write_text(
        json.dumps(extracted, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"  {len(extracted['slides'])} 页")

    print("[2/4] 生成（调模型）")
    key = read_key()
    material = generate(key, extracted)
    (out_dir / "material.json").write_text(
        json.dumps(material, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print("[3/4] 渲染 HTML")
    html_path = out_dir / "review.html"
    render(extracted, material, html_path)

    print("[4/4] 出 PDF")
    pdf_path = out_dir / "review.pdf"
    to_pdf(html_path, pdf_path)

    print(f"\n完成：\n  HTML: {html_path}\n  PDF : {pdf_path}")


if __name__ == "__main__":
    main()
